# Cost, cash and receivables from the Wafeq accounting exports Tarik shares in the first week
# of each month (after the monthly invoices are released):
#   sources/wafeq_bills_<date>.xlsx      — purchase bills  (فواتير المشتريات) → the expense side
#   sources/wafeq_journal_<date>.xlsx    — journal entries (القيود اليومية)   → salaries, funding, receipts
#   sources/ar_statement_<date>.json     — customer balances (كشف أرصدة العملاء), transcribed from the PDF
# The sales-invoice export in the same drop is only a cross-check: billing itself comes from the API.
import os, json, glob
from openpyxl import load_workbook

HERE = os.path.dirname(os.path.abspath(__file__))
SRC = os.path.join(HERE, 'sources')

# Bill accounts spent on a client's behalf (ads, prep, shipping, marketplace fees): rebill candidates,
# not Nasam operating cost. Every other bill account is treated as Nasam's own cost.
CLIENT_SIDE = {
    'الإعلانات داخل المنصات': 'Client ad spend inside the marketplaces',
    'تجهيز وتغليف وتخزين': 'Prep, packing and storage for client stock',
    'عمولة المنصات': 'Marketplace commission paid on client sales',
    'تكاليف الشحن': 'Shipping paid on client orders',
    'خصومات وعروض ترويجية': 'Promotional discounts funded on client stores',
    'تخزين': 'Storage',
    'تخزين تالف': 'Damaged-stock storage',
}
SALARY_CODES = {'5201', '5231'}           # staff salaries: operational + administrative
FUNDING_CODES = {'3231', '3232', '3233', '3234'}   # partner contributions (funding, not revenue)
GRANT_CODES = {'421'}                     # Monshaat support
SETTLE_CODES = {'418'}                    # revenue under settlement (Salla payouts land here)
ZAKAT_CODES = {'545'}


def _latest(pattern):
    f = sorted(glob.glob(os.path.join(SRC, pattern)))
    return f[-1] if f else None


def available():
    return bool((_latest('wafeq_bills_*.xlsx') and _latest('wafeq_journal_*.xlsx'))
                or _latest('acct_series_*.json'))


def as_of():
    """Date stamp of the accounting drop the cost sheet is built from."""
    f = _latest('wafeq_bills_*.xlsx')
    if f: return os.path.basename(f).replace('wafeq_bills_', '').replace('.xlsx', '')
    c = _latest('acct_series_*.json') or ''
    return os.path.basename(c).replace('acct_series_', '').replace('.json', '')


def _rows(path):
    ws = load_workbook(path, data_only=True).active
    it = ws.iter_rows(values_only=True)
    next(it)
    return [r for r in it if any(v is not None for v in r)]


def _f(v):
    return float(v) if isinstance(v, (int, float)) else 0.0


def bills():
    """One row per bill line: month, account, amount, contact, bill number, client-side flag."""
    out = []
    for r in _rows(_latest('wafeq_bills_*.xlsx')):
        acct = str(r[22] or '').strip()
        out.append(dict(num=r[1], date=str(r[2])[:10], month=str(r[2])[:7], contact=str(r[4] or ''),
                        desc=str(r[19] or ''), account=acct, amount=_f(r[26]),
                        client_side=acct in CLIENT_SIDE))
    return out


def journal():
    """One row per journal line: month, account code and name, debit, credit."""
    out = []
    for r in _rows(_latest('wafeq_journal_*.xlsx')):
        out.append(dict(date=str(r[3])[:10], month=str(r[3])[:7], desc=str(r[8] or ''),
                        account=str(r[9] or '').strip(), code=str(r[10] or '').strip(),
                        debit=_f(r[11]), credit=_f(r[12])))
    return out


def ar_statement():
    p = _latest('ar_statement_*.json')
    return json.load(open(p, encoding='utf-8')) if p else None


def series():
    """Monthly series for the cost sheet. Keys are 'YYYY-MM'.

    The raw bills and journal exports carry pay and insurance amounts against named staff, so they
    are NOT committed (see .gitignore). Whenever they are present this writes the aggregates it needs
    to sources/acct_series_<date>.json — account x month totals only, no names, no line detail — and
    that file is what the repo keeps and a fresh clone builds from.
    """
    from collections import defaultdict
    if not (_latest('wafeq_bills_*.xlsx') and _latest('wafeq_journal_*.xlsx')):
        cached=_latest('acct_series_*.json')
        if not cached: raise SystemExit('no accounting export and no cached series in sources/')
        d=json.load(open(cached, encoding='utf-8'))
        return {k: defaultdict(float, {(tuple(x[0]) if isinstance(x[0], list) else x[0]): x[1] for x in v})
                for k, v in d['series'].items()}
    op = defaultdict(float)          # Nasam operating cost, by (account, month)
    cs = defaultdict(float)          # client-side / rebill candidates, by (account, month)
    for b in bills():
        (cs if b['client_side'] else op)[(b['account'], b['month'])] += b['amount']
    sal = defaultdict(float)         # salaries by month (debit less credit reversals)
    fund = defaultdict(float)        # partner contributions by month
    grant = defaultdict(float)       # Monshaat support
    settle = defaultdict(float)      # Salla payouts received into settlement
    zakat = defaultdict(float)
    gosi_adj = defaultdict(float)    # GOSI / insurance reversals booked in the journal
    for j in journal():
        m, c = j['month'], j['code']
        if c in SALARY_CODES:
            sal[m] += j['debit'] - j['credit']
        elif c in FUNDING_CODES:
            fund[m] += j['credit'] - j['debit']
        elif c in GRANT_CODES:
            grant[m] += j['credit'] - j['debit']
        elif c in SETTLE_CODES:
            settle[m] += j['credit'] - j['debit']
        elif c in ZAKAT_CODES:
            zakat[m] += j['debit'] - j['credit']
        elif c in ('5202', '5241'):
            gosi_adj[m] += j['debit'] - j['credit']
    out=dict(operating=op, client_side=cs, salaries=sal, funding=fund,
             grant=grant, settlement=settle, zakat=zakat, staff_adj=gosi_adj)
    dest=os.path.join(SRC, f'acct_series_{as_of()}.json')
    json.dump({'as_of': as_of(),
               'note': 'Aggregates generated by accounting.py from the Wafeq bills and journal exports. '
                       'Account x month totals only — the raw exports are not committed.',
               'series': {k: [[list(kk) if isinstance(kk, tuple) else kk, round(vv, 2)] for kk, vv in v.items()]
                          for k, v in out.items()}},
              open(dest, 'w', encoding='utf-8'), ensure_ascii=False, indent=1)
    return out
